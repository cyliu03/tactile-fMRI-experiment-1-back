"""Shared, file-independent schedule design for the visual and tactile tasks.

The CSV is the object inventory, never a directory scan. The Python generators
use openpyxl so they can run on the experiment PC with scripts/requirements.txt.
"""
from __future__ import annotations

import argparse
import csv
import json
import random
from collections import Counter, defaultdict
from itertools import combinations
from pathlib import Path, PurePosixPath

ROOT = Path(__file__).resolve().parents[1]
CONFIG_DIR = ROOT / "config"
TRIAL_COLUMNS = [
    "session", "session_run", "run", "trial", "trial_id", "stimulus_number",
    "plate", "plate_row", "plate_col",
    "base_trial", "nback_target", "correct_response", "question_type",
    "question_text", "iti_s",
]
RUN_COLUMNS = [
    "session", "session_run", "run", "plate", "order", "n_base_trials",
    "nback_targets", "n_trials", "iti_sum_s", "duration_s",
]
CATALOG_COLUMNS = ["stimulus_number", "object_name", "video_filename"]


def require(condition, message):
    if not condition:
        raise ValueError(message)


def load_inputs(catalog_path, config_path):
    config = json.loads(Path(config_path).read_text(encoding="utf-8-sig"))
    with Path(catalog_path).open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle)
        require("stimulus_number" in (reader.fieldnames or []), "物体清单需要 stimulus_number 列")
        catalog = list(reader)
    require(bool(catalog), "物体清单不能为空")
    ids, videos = set(), set()
    groups = defaultdict(list)
    for item in catalog:
        for key in CATALOG_COLUMNS:
            item[key] = (item.get(key) or "").strip()
        number = int(item["stimulus_number"])
        scheme = config.get("numbering_scheme", "session_hundreds")
        if scheme == "session_hundreds":
            session = number // 100
            require(session > 0 and 1 <= number - session * 100 <= 99,
                    f"编号 {number} 必须采用 session×100＋1～99 的命名规则")
        else:
            require(scheme == "sequential" and 1 <= number <= 99, "test 编号必须为 001–099")
            session = (number - 1) // config["objects_per_session"] + 1
        require(number not in ids, f"物体编号重复: {number}")
        name = item["video_filename"] or f"{number:03d}.mp4"
        require("/" not in name and "\\" not in name and ":" not in name and
                PurePosixPath(name).suffix.lower() == ".mp4",
                f"video_filename 只填写 .mp4 文件名: {name}")
        require(name not in videos, f"视频文件名重复: {name}")
        videos.add(name)
        ids.add(number)
        item["session"], item["stimulus_number"] = session, number
        item = {key: item[key] for key in ["session", *CATALOG_COLUMNS]}
        groups[session].append(item)
    for key in ("seed", "base_trials_per_run", "nback_targets_per_run", "min_nback_target_gap", "plate_rows", "plate_cols"):
        require(type(config[key]) is int and config[key] >= (0 if key == "seed" else 1), f"无效参数: {key}")
    base = config["base_trials_per_run"]
    require(base == config["plate_rows"] * config["plate_cols"], "每 run 基础 trial 数必须等于盘的行×列")
    require(config["run_orders"] and all(x in {"ascending", "descending", "column_major", "reverse_column_major"}
            for x in config["run_orders"]), "run_orders 中有未知排序方式")
    require(type(config.get("runs_per_session", 4)) is int and config.get("runs_per_session", 4) > 0,
            "runs_per_session 必须为正整数")
    for session, items in groups.items():
        items.sort(key=lambda x: x["stimulus_number"])
        require(len(items) % base == 0, f"Session {session} 有 {len(items)} 个物体，必须为 {base} 的整倍数（完整盘）")
    timing = config["timing"]
    for key in ("delivery_s", "prepare_s", "touch_s", "hold_s", "video_s", "post_video_hold_s", "response_s", "initial_dummy_s", "final_rest_s"):
        require(isinstance(timing[key], (int, float)) and timing[key] >= 0, f"无效时间: {key}")
    options = timing["iti_options_s"]
    require(options and all(isinstance(x, (int, float)) and x > 0 for x in options), "ITI 选项必须是正数")
    require(sum(options) / len(options) == timing["iti_mean_per_run_s"], "ITI 选项均值不匹配")
    require(bool(config["video_root"]) and not PurePosixPath(config["video_root"]).is_absolute()
            and ":" not in config["video_root"], "video_root 使用相对程序目录的路径")
    return config, dict(sorted(groups.items()))


def iti_pool(config, rng):
    t = config["timing"]
    count = config["base_trials_per_run"] + config["nback_targets_per_run"]
    times, extra = divmod(count, len(t["iti_options_s"]))
    pool = list(t["iti_options_s"]) * times
    extras = sorted(t["iti_options_s"], key=lambda x: (abs(x - t["iti_mean_per_run_s"]), x))[:extra]
    require(sum(extras) == extra * t["iti_mean_per_run_s"], "当前 trial 数不能构造指定均值的 ITI")
    pool.extend(extras)
    rng.shuffle(pool)
    return pool


def ordered_plate(plate, order, config):
    if order == "ascending":
        return list(plate)
    if order == "descending":
        return list(reversed(plate))
    cols = range(config["plate_cols"])
    if order == "reverse_column_major":
        cols = reversed(list(cols))
    return [plate[row * config["plate_cols"] + col]
            for col in cols for row in range(config["plate_rows"])]


def build_schedule(modality, config, groups):
    require(modality in {"tactile", "visual"}, "未知模态")
    target_rng = random.Random(config["seed"])
    iti_rng = random.Random(config["seed"] + 1)
    used_targets, position_counts = set(), Counter()
    trials, runs, inventory = [], [], []
    base = config["base_trials_per_run"]
    t = config["timing"]
    fixed_s = t["prepare_s"] + t["response_s"]
    fixed_s += (t["video_s"] + t["post_video_hold_s"] if modality == "visual"
                else t["delivery_s"] + t["touch_s"] + t["hold_s"])
    for session, items in groups.items():
        for index, item in enumerate(items):
            item = dict(item)
            item["plate"] = index // base + 1
            item["plate_row"] = index % base // config["plate_cols"] + 1
            item["plate_col"] = index % config["plate_cols"] + 1
            inventory.append(item)
        session_items = [x for x in inventory if x["session"] == session]
        session_run = 0
        for order in config["run_orders"]:
            for offset in range(0, len(session_items), base):
                session_run += 1
                run = len(runs) + 1
                plate = ordered_plate(session_items[offset:offset + base], order, config)
                choices = []
                available = [i for i, x in enumerate(plate) if x["stimulus_number"] not in used_targets]
                for indices in combinations(available, config["nback_targets_per_run"]):
                    positions = [i + rank + 2 for rank, i in enumerate(indices)]
                    if all(b - a >= config["min_nback_target_gap"] for a, b in zip(positions, positions[1:])):
                        choices.append((sum(position_counts[p] for p in positions), indices, positions))
                require(bool(choices), f"Session {session} / Run {run}: 无法满足 1-back 间距和不重复选取规则，请调整配置")
                penalty = min(x[0] for x in choices)
                _, selected, positions = target_rng.choice([x for x in choices if x[0] == penalty])
                used_targets.update(plate[i]["stimulus_number"] for i in selected)
                position_counts.update(positions)
                expanded = []
                for index, item in enumerate(plate):
                    expanded.append((item, index + 1, 0))
                    if index in selected:
                        expanded.append((item, index + 1, 1))
                itis = iti_pool(config, iti_rng)
                for trial_num, ((item, base_trial, target), iti) in enumerate(zip(expanded, itis), 1):
                    number = item["stimulus_number"]
                    trial = dict(zip(TRIAL_COLUMNS, [
                        session, session_run, run, trial_num, f"{len(trials) + 1:06d}", number,
                        item["plate"], item["plate_row"], item["plate_col"], base_trial, target,
                        "yes" if target else "no", "one_back_same", "是否和前一个物体一样？", iti,
                    ]))
                    if modality == "visual":
                        trial["video_filename"] = item["video_filename"]
                    trials.append(trial)
                runs.append(dict(zip(RUN_COLUMNS, [session, session_run, run, offset // base + 1,
                    order, base, len(selected), len(expanded), sum(itis),
                    len(expanded) * fixed_s + sum(itis) + t["initial_dummy_s"] + t["final_rest_s"]])))
    # Regroup after generating the configured full sequence, preserving its RNG
    # choices, global run/trial identities, ITIs and source stimulus folders.
    runs_per_session = config.get("runs_per_session", 4)
    for row in [*runs, *trials]:
        row["stimulus_session"] = row["session"]
        row["session"] = (row["run"] - 1) // runs_per_session + 1
        row["session_run"] = (row["run"] - 1) % runs_per_session + 1
    params = {
        "schema_version": 4, "modality": modality, "seed": config["seed"],
        "dataset": config.get("dataset", "train"),
        "numbering_scheme": config.get("numbering_scheme", "session_hundreds"),
        "objects_per_session": config.get("objects_per_session", 60),
        "runs_per_session": runs_per_session,
        "n_stimulus_sessions": len(groups),
        "video_session_pattern": "session%02d",
        "n_objects": len(inventory), "n_sessions": (len(runs) + runs_per_session - 1) // runs_per_session, "n_runs": len(runs),
        "n_base_trials_per_run": base, "nback_targets_per_run": config["nback_targets_per_run"],
        "n_trials_per_run": base + config["nback_targets_per_run"],
        "n_total_trials": len(trials), "n_base_trials": base * len(runs),
        "n_total_nback_targets": config["nback_targets_per_run"] * len(runs),
        "base_repetitions_per_object": len(config["run_orders"]),
        "min_nback_target_gap": config["min_nback_target_gap"],
        "video_root": config["video_root"] if modality == "visual" else "n/a",
        "trigger_key": ",".join(config.get("trigger_keys", ["s"])), "response_hand": "右手", "touch_hand": "左手" if modality == "tactile" else "n/a",
        **{key: value for key, value in t.items() if key != "iti_options_s"},
        "iti_options_s": ",".join(map(str, t["iti_options_s"])),
    }
    return {"parameters": params, "runs": runs, "trials": trials, "stimuli": inventory}


def workbook_specs(schedule, modality, config):
    p, runs, trials = schedule["parameters"], schedule["runs"], schedule["trials"]
    overview = [["项目", "内容"], ["模态", "视觉" if modality == "visual" else "触觉"],
        ["Session 数", p["n_sessions"]], ["全局 Run 数", p["n_runs"]],
        ["每 Run 基础 Trial", p["n_base_trials_per_run"]], ["每 Run 1-back 重复", p["nback_targets_per_run"]],
        ["每 Run 实际 Trial", p["n_trials_per_run"]], ["实际 Trial 总数", p["n_total_trials"]],
        ["独立物体数", p["n_objects"]], ["每物体基础呈现次数（本模态）", p["base_repetitions_per_object"]],
        ["编号规则", "101–160、201–260、301–360、401–460、501–560、601–660；百位及以上为 session"],
        ["执行顺序", "Trials 为实际执行表；按 session、全局 run、trial 排序；每次启动执行一个 run"],
        ["Run 编号", "全局 run 连续编号；session_run 表示该 session 内的轮次；输入框填写全局 run"],
        ["基础排序", "每 session 两盘：升序盘1/盘2，降序盘1/盘2，逐列盘1/盘2；盘为5行6列"],
        ["1-back", "重复紧跟原 trial；target=1 的正确答案为 yes；每 run 首 trial 为 no；重复不跨 run"],
        ["随机规则", f"seed={p['seed']}；target 间隔至少 {p['min_nback_target_gap']} trial；同一物体本模态最多选为 target 一次"],
        ["ITI", "seed+1 随机打乱 2/3/4/5/6 秒；每 run ITI 均值4秒，合计132秒"],
        ["时长", "根据阶段参数与逐trial的ITI计算"],
        ["物体/视频未齐", "生成不检查文件存在；视觉正式执行前只检查当前 run 视频；触觉不打开 STL"],
        ["唯一物体标识", "只维护 stimulus_number；object_name 可留空；视觉 video_filename 留空则自动使用 编号.mp4"],
        ["视频位置", config["video_root"] if modality == "visual" else "触觉无物体路径依赖"],
        ["新增物体", "编辑 config/stimulus_catalog.csv；完整session按60个编号追加；已有session按30个的完整盘追加"],
        ["配置", "config/experiment_design.json 控制盘大小、排序、重复数、时长；修改后重新生成"],
        ["MATLAB数据源", "仅同目录 experiment_schedule_seed42.xlsx；Parameters/Runs/Trials 为机器读取表"],
        ["两模态对应", "两模态使用同一stimulus_number；分析输出的ID自动生成；当前seed相同时顺序、target及ITI相同"],
    ]
    # Keep descriptive values correct for future custom designs, too.
    dataset = p["dataset"]
    overview[10][1] = ("百位及以上为刺激来源组；" if dataset == "train" else "连续编号，每30个为刺激来源组；") + "原目录：" + "；".join(
        f"S{s}: {min(x['stimulus_number'] for x in schedule['stimuli'] if x['session'] == s):03d}–"
        f"{max(x['stimulus_number'] for x in schedule['stimuli'] if x['session'] == s):03d}"
        for s in sorted({x['session'] for x in schedule['stimuli']}))
    execution_note = f"每个测试物体基础呈现{p['base_repetitions_per_object']}次" if dataset == "test" else "保留全部原run"
    overview[11][1] = f"{execution_note}；每{p['runs_per_session']}个全局run组成一个输入Session；最后一组可不足{p['runs_per_session']}个"
    overview[12][1] = "输入框填写 session 和全局 run；session_run 为组内轮次；stimulus_session 为原刺激目录编号"
    overview[13][1] = "排序：" + ", ".join(config["run_orders"]) + f"；依次遍历各盘；每盘{config['plate_rows']}行×{config['plate_cols']}列"
    overview[16][1] = f"seed+1 打乱 {p['iti_options_s']} 秒；每run均值{p['iti_mean_per_run_s']}秒"
    run_seconds = "、".join(str(x) for x in sorted({r['duration_s'] for r in runs}))
    overview[17][1] = f"每run {run_seconds}秒；开始基线{p['initial_dummy_s']}秒、结束静息{p['final_rest_s']}秒；不含触发前等待和人工换run"
    overview[19][1] = "只维护 stimulus_number；默认视频名为三位编号.mp4（如001.mp4）；名称及特殊视频名可留空"
    overview[20][1] = (config["video_root"] + "/sessionXX/编号.mp4；XX 使用 stimulus_session，原目录不变" if modality == "visual"
                        else "触觉仅使用编号；原 object/" + dataset + "/sessionXX 目录保留，XX 对应刺激来源组")
    overview[21][1] = "编辑主目录下 config/" + ("stimulus_catalog.csv" if dataset == "train" else "stimulus_catalog_test.csv") + "；按完整盘追加，每盘30个物体"
    overview[22][1] = "主目录下 config/" + ("experiment_design.json" if dataset == "train" else "experiment_design_test.json") + " 控制排序、重复数、时长和视频根目录"
    overview[23][1] = "仅同目录 " + config.get("schedule_filename", "experiment_schedule_seed42_current.xlsx") + "；Parameters/Runs/Trials 为机器读取表；旧表不参与执行"
    overview.append(["数据集", dataset])
    trial_columns = TRIAL_COLUMNS + (["video_filename"] if modality == "visual" else []) + ["stimulus_session"]
    run_columns = RUN_COLUMNS + ["stimulus_session"]
    stimulus_columns = ["stimulus_number", "object_name"] + (["video_filename"] if modality == "visual" else [])
    specs = [{"name": "Overview", "rows": overview, "widths": [36, 113]},
        {"name": "Runs", "rows": [run_columns] + [[r[k] for k in run_columns] for r in runs], "widths": [12, 16, 12, 12, 23, 18, 18, 14, 16, 16, 20]},
        {"name": "Trials", "rows": [trial_columns] + [[r[k] for k in trial_columns] for r in trials], "widths": [12,16,10,10,14,22,10,14,14,16,18,22,22,30,12] + ([26] if modality == "visual" else []) + [20], "target_column": trial_columns.index("nback_target"), "number_columns": [trial_columns.index("stimulus_number")]},
        {"name": "Stimuli", "rows": [stimulus_columns] + [[r[k] for k in stimulus_columns] for r in schedule["stimuli"]], "widths": [22,30,30], "number_columns": [0]},
        {"name": "Parameters", "rows": [["key", "value"]] + [[k,v] for k,v in p.items()], "widths": [38,45]},
    ]
    printable = []
    for session in sorted({x["session"] for x in runs}):
        current = [r for r in runs if r["session"] == session]
        columns = [[x for x in trials if x["run"] == r["run"]] for r in current]
        rows = [["Trial"] + [f"Run {r['run']} / S{session}-R{r['session_run']}" for r in current]]
        for i in range(p["n_trials_per_run"]):
            rows.append([i+1] + [f"{c[i]['stimulus_number']:03d}" + (" [1-back]" if c[i]["nback_target"] else "") for c in columns])
        printable.append({"name": f"Session {session:02d}", "rows": rows, "widths": [9] + [25]*len(current)})
    return specs, printable


def write_openpyxl(specs, output):
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as exc:
        raise RuntimeError("请在主目录执行 python -m pip install -r scripts/requirements.txt") from exc
    wb = Workbook()
    wb.remove(wb.active)
    for spec in specs:
        ws = wb.create_sheet(spec["name"])
        ws.sheet_view.showGridLines = False
        for row in spec["rows"]:
            ws.append(row)
        for row in ws:
            for cell in row:
                cell.font = Font(name="Microsoft YaHei", size=10)
                cell.alignment = Alignment(vertical="center", wrap_text=True)
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="24476A")
            cell.font = Font(name="Microsoft YaHei", size=10, bold=True, color="FFFFFF")
        for i, width in enumerate(spec["widths"], 1):
            ws.column_dimensions[get_column_letter(i)].width = width
        for col in spec.get("number_columns", []):
            for cells in ws.iter_rows(min_row=2, min_col=col+1, max_col=col+1):
                cells[0].number_format = "000"
        for i in range(1, ws.max_row+1):
            ws.row_dimensions[i].height = 30 if spec["name"] == "Overview" else 22
        ws.freeze_panes = "B2" if spec["name"].startswith("Session") else "A2"
        if "target_column" in spec:
            for row in ws.iter_rows(min_row=2):
                if row[spec["target_column"]].value == 1:
                    for cell in row:
                        cell.fill = PatternFill("solid", fgColor="FFF0C2")
        if spec["name"].startswith("Session"):
            for row in ws.iter_rows(min_row=2):
                for cell in row:
                    if "[1-back]" in str(cell.value):
                        cell.fill = PatternFill("solid", fgColor="FFF0C2")
            ws.page_setup.orientation = "landscape"
            ws.page_setup.paperSize = ws.PAPERSIZE_A4
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 1
            ws.sheet_properties.pageSetUpPr.fitToPage = True
            ws.print_title_rows = "1:1"
        elif spec["name"] in {"Runs", "Trials", "Stimuli"}:
            ws.auto_filter.ref = ws.dimensions
    # Never silently change the execution filename when Excel has locked it.
    wb.save(output)


def main(modality, script_dir, dataset="train"):
    parser = argparse.ArgumentParser(description="从物体清单生成Excel执行表，不检查STL或视频是否存在")
    suffix = "" if dataset == "train" else "_test"
    parser.add_argument("--catalog", type=Path, default=CONFIG_DIR / f"stimulus_catalog{suffix}.csv")
    parser.add_argument("--config", type=Path, default=CONFIG_DIR / f"experiment_design{suffix}.json")
    parser.add_argument("--output-dir", type=Path, default=Path(script_dir))
    parser.add_argument("--spec-only", action="store_true", help=argparse.SUPPRESS)
    args = parser.parse_args()
    config, groups = load_inputs(args.catalog, args.config)
    require(config.get("dataset", "train") == dataset, "清单配置数据集与程序目录不一致")
    schedule = build_schedule(modality, config, groups)
    specs, printable = workbook_specs(schedule, modality, config)
    schedule_filename = config.get("schedule_filename", f"experiment_schedule_seed{config['seed']}_current.xlsx")
    args.output_dir.mkdir(parents=True, exist_ok=True)
    if args.spec_only:
        path = args.output_dir / f"{dataset}_{modality}_workbook_spec.json"
        path.write_text(json.dumps({"schedule": schedule, "workbooks": [
            {"filename": schedule_filename, "sheets": specs},
            {"filename": f"printable_run_order_seed{config['seed']}.xlsx", "sheets": printable},
        ]}, ensure_ascii=False), encoding="utf-8")
    else:
        for filename, sheets in [(schedule_filename, specs),
                                  (f"printable_run_order_seed{config['seed']}.xlsx", printable)]:
            write_openpyxl(sheets, args.output_dir / filename)
            print(args.output_dir / filename)
    p = schedule["parameters"]
    print(f"{modality}: {p['n_objects']} objects, {p['n_sessions']} sessions, {p['n_runs']} runs, {p['n_total_trials']} trials")
