"""Read-only audit of the four current schedules, paths and asset migration."""
from pathlib import Path
import hashlib
import json
import re
from openpyxl import load_workbook
from schedule_core import CONFIG_DIR, load_inputs, build_schedule, workbook_specs

ROOT = Path(__file__).resolve().parents[1]
normalize = lambda x: None if x == '' else x


def audit():
    results = []
    for dataset in ['train','test']:
        suffix = '' if dataset == 'train' else '_test'
        config, groups = load_inputs(CONFIG_DIR/f'stimulus_catalog{suffix}.csv',CONFIG_DIR/f'experiment_design{suffix}.json')
        for modality in ['visual','tactile']:
            folder = ROOT/f'{dataset}_{modality}'
            current = folder/config['schedule_filename']
            wb = load_workbook(current,data_only=False)
            schedule = build_schedule(modality,config,groups)
            specs, printable = workbook_specs(schedule,modality,config)
            for spec in specs:
                ws = wb[spec['name']]
                expected = [[normalize(v) for v in row] for row in spec['rows']]
                nc,nr = len(expected[0]),len(expected)
                actual = [list(r[:nc]) for r in ws.values][:nr]
                assert actual == expected, (folder.name,ws.title,'saved values do not match catalog/config')
                assert all(all(v is None for v in row[nc:]) for row in ws.values), 'Unexpected old columns'
                for col in spec.get('number_columns',[]):
                    assert ws.cell(2,col+1).number_format == '000', 'Three-digit object display was lost'
            params = dict(wb['Parameters'].iter_rows(min_row=2,values_only=True))
            assert params['schema_version'] == 4 and params['dataset'] == dataset
            assert params['runs_per_session'] == 4
            if modality == 'visual':
                assert params['video_root'] == f'../video/{dataset}'
                assert params['video_session_pattern'] == 'session%02d'
            else:
                assert params['video_root'] == 'n/a'
                assert 'video_filename' not in list(wb['Trials'].values)[0]
            source = (folder/f'{modality}_fMRI_experiment.m').read_text(encoding='utf-8')
            assert "'experiment_schedule_seed42_current.xlsx'" in source
            for forbidden in ['Standard_Objects','data/real_shapes','experiment_schedule_*.json','stimulus_order_number']:
                assert forbidden not in source, (folder,forbidden)
            assert not any(old in list(wb['Trials'].values)[0] for old in ['stimulus_id','object_id','stimulus_name','object_filename','stim_file'])
            flat_strings = [value for ws in wb for row in ws.values for value in row if isinstance(value,str)]
            assert not any('Standard_Objects' in v or 'data/real_shapes' in v for v in flat_strings)
            print_wb=load_workbook(folder/'printable_run_order_seed42.xlsx',data_only=True)
            assert print_wb.sheetnames == [spec['name'] for spec in printable]
            for spec in printable:
                width=len(spec['rows'][0])
                assert [list(r[:width]) for r in print_wb[spec['name']].values] == spec['rows']
                assert all(all(v is None for v in r[width:]) for r in print_wb[spec['name']].values)
            missing,present = [],[]
            for item in schedule['stimuli']:
                name = item['video_filename'] or f"{item['stimulus_number']:03d}.mp4"
                expected_path = ROOT/'video'/dataset/f"session{item['session']:02d}"/name
                assert expected_path.parent.is_dir(), expected_path.parent
                assert (ROOT/'object'/dataset/f"session{item['session']:02d}").is_dir()
                if modality == 'visual':
                    exists=expected_path.is_file()
                    label=expected_path.relative_to(ROOT).as_posix()
                else:
                    object_folder=ROOT/'object'/dataset/f"session{item['session']:02d}"
                    exists=any(int(re.match(r'^(\d+)',p.name).group(1))==item['stimulus_number']
                               for p in object_folder.glob('*.stl') if re.match(r'^(\d+)',p.name))
                    label=(object_folder/f"{item['stimulus_number']:03d}*.stl").relative_to(ROOT).as_posix()
                (present if exists else missing).append(label)
            p=schedule['parameters']
            result={'program':folder.name,'objects':p['n_objects'],'sessions':p['n_sessions'],
                    'runs':p['n_runs'],'base_trials':p['n_base_trials'],'actual_trials':p['n_total_trials'],
                    'asset_type':'video' if modality=='visual' else 'object (inventory only, not a runtime dependency)',
                    'assets_present':len(present),'assets_not_yet_supplied':len(missing),'missing_assets':missing}
            results.append(result)
            print(f"PASS {folder.name}: {p['n_sessions']} sessions / {p['n_runs']} runs / {p['n_total_trials']} actual trials; {result['asset_type']} {len(present)}/{p['n_objects']}")
    manifest = ROOT/'schedule_backup/schedule_backup_20261017_layout/authoring/asset_migration.json'
    migrations = json.loads(manifest.read_text(encoding='utf-8'))
    for item in migrations:
        path=ROOT/item['new_path']
        assert path.is_file() and hashlib.sha256(path.read_bytes()).hexdigest()==item['sha256'], path
    assert len(migrations)==120
    assert not (ROOT/'Standard_Objects').exists()
    assert not list((ROOT/'video').glob('*.mp4'))
    assert not list((ROOT/'object').glob('*.stl'))
    print(f'PASS media migration: {len(migrations)} file hashes preserved')
    return results


if __name__ == '__main__':
    audit()
